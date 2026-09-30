"""Admin tools: order filters (+ saved filters, Excel export) and bulk upload (orders, members) with undo."""
import io
import uuid
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient

PW = "Pass-12345"
TODAY = date.today()


@pytest.fixture()
def client(migrated_db_url, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", migrated_db_url)
    monkeypatch.setenv("JWT_SECRET", "x" * 48)
    from app import config
    from app.main import app
    config._view_cache = None
    config._perm_cache = None
    with TestClient(app) as c:
        yield c


@pytest.fixture()
def seed(db_conn):
    with db_conn.cursor() as cur:
        for i in range(-40, 3):
            d = TODAY + timedelta(days=i)
            cur.execute("INSERT INTO dim_date (date_key, full_date, day, month, year, day_of_week, is_monday_holiday) "
                        "VALUES (%s,%s,%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING",
                        (int(d.strftime("%Y%m%d")), d, d.day, d.month, d.year, d.strftime("%A"), d.weekday() == 0))
        for table, col, names in (("dim_order_channel", "channel_name", ("Call", "Walk-in", "Online")),
                                  ("dim_submission_type", "type_name", ("Challan", "Invoice")),
                                  ("dim_delivery_status", "status_name", ("Shop", "Delivered", "Cancelled")),
                                  ("dim_payment_status", "status_name", ("Paid", "Pending"))):
            for n in names:
                cur.execute(f"INSERT INTO {table} ({col}) VALUES (%s) ON CONFLICT DO NOTHING", (n,))
        for n, r in (("Ravi", "ready_by"), ("Sonu", "colour_making"), ("Amit", "delivery")):
            cur.execute("INSERT INTO dim_person (full_name, person_role) VALUES (%s,%s) ON CONFLICT DO NOTHING", (n, r))
    db_conn.commit()


def login(client, db_conn, role="admin"):
    from app import auth
    username = f"bt_{role}_{uuid.uuid4().hex[:6]}"
    with db_conn.cursor() as cur:
        cur.execute("INSERT INTO dim_user (username, password_hash, role, display_name, must_change_password) "
                    "VALUES (%s,%s,%s,%s,false)", (username, auth.hash_password(PW), role, "BT " + role))
    db_conn.commit()
    tok = client.post("/auth/login", json={"username": username, "password": PW}).json()["access_token"]
    return {"Authorization": f"Bearer {tok}"}


HEAD = ("Order Date,Order Via,Type of Submission,DC/Inv No,Address,Remarks,Ready By,Colour Making By,Delivery Status,"
        "Delivered By,Delivery Date & Time,Cartage,Date of Receiving,Payment Status,Amount Received")


def d(offset):
    return (TODAY + timedelta(days=offset)).strftime("%d-%m-%Y")


def order_csv(*lines, head=HEAD):
    return ("\n".join([head, *lines]) + "\n").encode("utf-8")


def upload(client, h, entity, raw, name="f.csv"):
    r = client.post(f"/admin/bulk/{entity}/validate?filename={name}", content=raw, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def confirm(client, h, entity, rows, name="f.csv"):
    r = client.post(f"/admin/bulk/{entity}/confirm", json={"rows": rows, "filename": name}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


# ---------------------------------------------------------------- template + access
def test_template_and_admin_only(client, db_conn, seed):
    admin, shop = login(client, db_conn), login(client, db_conn, "shop")
    r = client.get("/admin/bulk/orders/template.csv", headers=admin)
    assert r.status_code == 200 and r.text.lstrip("﻿").startswith("Order Date,Order Via,Type of Submission")
    assert "EXAMPLE" in r.text
    assert client.get("/admin/bulk/orders/template.csv", headers=shop).status_code == 403
    assert client.get("/admin/bulk/nonsense/template.csv", headers=admin).status_code == 404
    assert client.post("/admin/bulk/orders/validate", content=b"x", headers=shop).status_code == 403
    assert client.get("/admin/orders/search", headers=shop).status_code == 403


def test_template_round_trips_and_examples_are_ignored(client, db_conn, seed):
    admin = login(client, db_conn)
    tpl = client.get("/admin/bulk/orders/template.csv", headers=admin).content
    rep = upload(client, admin, "orders", tpl)
    assert "format_error" in rep and "No data rows" in rep["format_error"]


# ---------------------------------------------------------------- order validation
def test_valid_rows_and_all_the_error_messages(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    raw = order_csv(
        f"{d(-10)},Call,Challan,V-{u}-1,Rohini,note,Ravi,Sonu,Delivered,Amit,{d(-10)} 14:30,50,{d(-9)},Paid,1500",
        f"{d(-10)},Walk-in,Invoice,V-{u}-2,,,,,,,,,,,",
        f"31-02-2026,Call,Challan,V-{u}-3,,,,,,,,,,,",                       # not a date
        f"{d(-10)},Phone,Challan,V-{u}-4,,,,,,,,,,,",                         # unknown channel
        f"{d(-10)},Call,Challan,,,,,,,,,,,,",                                 # missing DC
        f"{d(-10)},Call,Challan,V-{u}-1,,,,,,,,,,,",                          # duplicate of row 2 in the file
        f"{d(-10)},Call,Challan,V-{u}-6,,,,,,,,,{d(-9)},,",                   # received without delivery
        f"{d(-10)},Call,Challan,V-{u}-7,,,,,,,,-5,,,",                        # negative cartage
        f"{d(5)},Call,Challan,V-{u}-8,,,,,,,,,,,",                            # future
        f"{d(-10)},Call,Challan,V-{u}-9,,,Nobody,,,,,,,,",                    # unknown person
        f"{d(-10)},Call,Challan,V-{u}-10,,,,,,,{d(-11)} 10:00,,,,",           # delivered before ordered
        "EXAMPLE 01-01-2026,Call,Challan,EXAMPLE-1,,,,,,,,,,,",               # ignored
        head=HEAD)
    rep = upload(client, admin, "orders", raw)
    assert rep["total"] == 11 and rep["valid"] == 2
    by_row = {e["row"]: e["error"] for e in rep["errors"]}
    assert "not a date" in by_row[4]
    assert "'Phone' not found" in by_row[5] and "Call" in by_row[5]          # lists the valid values
    assert "DC/Inv No is required" in by_row[6]
    assert "Duplicate of row 2" in by_row[7]
    assert "needs a Delivery Date" in by_row[8]
    assert "Cartage" in by_row[9]
    assert "future" in by_row[10]
    assert "'Nobody' not found" in by_row[11]
    assert "before the Order Date" in by_row[12]
    assert all(e["data"]["_row"] == e["row"] for e in rep["errors"])


def test_wrong_file_and_excel_upload(client, db_conn, seed):
    admin = login(client, db_conn)
    bad = upload(client, admin, "orders", b"name,age\nbob,3\n")
    assert "doesn't match the template" in bad["format_error"]
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD.split(","))
    ws.append([TODAY - timedelta(days=12), "Call", "Challan", "X-" + uuid.uuid4().hex[:5], "Rohini", "", "", "", "",
               "", "", 25, "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    rep = upload(client, admin, "orders", buf.getvalue(), "f.xlsx")
    assert rep["total"] == 1 and rep["valid"] == 1, rep["errors"]


def test_existing_order_is_flagged_and_row_cap(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    raw = order_csv(f"{d(-8)},Call,Challan,DB-{u},,,,,,,,,,,")
    rep = upload(client, admin, "orders", raw)
    confirm(client, admin, "orders", rep["rows"])
    again = upload(client, admin, "orders", raw)
    assert again["valid"] == 0 and "already exists" in again["errors"][0]["error"]
    many = order_csv(*[f"{d(-8)},Call,Challan,BIG-{i},,,,,,,,,,," for i in range(2001)])
    assert client.post("/admin/bulk/orders/validate", content=many, headers=admin).status_code == 400


def test_revalidate_lets_you_fix_a_row_without_reuploading(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    rep = upload(client, admin, "orders", order_csv(f"{d(-9)},Phone,Challan,FX-{u},,,,,,,,,,,"))
    assert rep["valid"] == 0
    row = rep["errors"][0]["data"]
    row["Order Via"] = "Call"
    fixed = client.post("/admin/bulk/orders/revalidate", json={"rows": [row]}, headers=admin).json()
    assert fixed["valid"] == 1 and not fixed["errors"]


# ---------------------------------------------------------------- confirm + undo
def test_import_creates_orders_with_batch_and_can_be_undone(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    raw = order_csv(
        f"{d(-20)},Call,Challan,IM-{u}-1,Rohini,,Ravi,Sonu,Delivered,Amit,{d(-20)} 14:30,50,{d(-19)},Paid,1,500.50",
        f"{d(-20)},Walk-in,Invoice,IM-{u}-2,Pitampura,,,,,,,,,,")
    raw = raw.replace(b"1,500.50", b'"1,500.50"')                         # a money value with a thousands comma
    rep = upload(client, admin, "orders", raw)
    assert rep["valid"] == 2, rep["errors"]
    res = confirm(client, admin, "orders", rep["rows"], "may.csv")
    assert res["created"] == 2
    batch = res["batch_id"]
    found = client.get(f"/admin/orders/search?batch_id={batch}&sort=sl_no&dir=asc", headers=admin).json()
    assert found["total"] == 2
    first, second = found["rows"]
    assert first["dc_inv_no"] == f"IM-{u}-1" and second["sl_no"] == first["sl_no"] + 1
    assert first["stage"] == "Closed" and float(first["amount_received"]) == 1500.50
    assert second["stage"] == "Awaiting godown"
    with db_conn.cursor() as cur:
        cur.execute("SELECT (timestamp_created AT TIME ZONE 'Asia/Kolkata')::time AS t, created_by_user_key, "
                    "import_batch_id FROM fact_orders WHERE sl_no = %s", (first["sl_no"],))
        row = cur.fetchone()
    assert str(row[0]) == "09:30:00" and row[2] == batch
    history = client.get("/admin/bulk/batches", headers=admin).json()
    assert history[0]["batch_id"] == batch and history[0]["filename"] == "may.csv" and history[0]["row_count"] == 2

    # importing the same rows again is refused, and nothing is written
    again = client.post("/admin/bulk/orders/confirm", json={"rows": rep["rows"]}, headers=admin).json()
    assert again["created"] == 0 and again["errors"]

    undo = client.post(f"/admin/bulk/batches/{batch}/undo", headers=admin)
    assert undo.status_code == 200 and undo.json()["removed"] == 2
    assert client.get(f"/admin/orders/search?batch_id={batch}", headers=admin).json()["total"] == 0
    assert client.post(f"/admin/bulk/batches/{batch}/undo", headers=admin).status_code == 409


def test_undo_is_refused_once_an_imported_order_was_edited(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    rep = upload(client, admin, "orders", order_csv(f"{d(-7)},Call,Challan,ED-{u},,,,,,,,,,,"))
    res = confirm(client, admin, "orders", rep["rows"])
    sl = client.get(f"/admin/orders/search?batch_id={res['batch_id']}", headers=admin).json()["rows"][0]["sl_no"]
    edit = client.put(f"/o2d/orders/{sl}/admin", json={"shippingLocation": "Changed"}, headers=admin).json()
    assert edit["success"] is True
    undo = client.post(f"/admin/bulk/batches/{res['batch_id']}/undo", headers=admin)
    assert undo.status_code == 409 and "edited" in undo.json()["detail"]
    assert client.get(f"/admin/orders/search?batch_id={res['batch_id']}", headers=admin).json()["total"] == 1


def test_confirm_is_all_or_nothing(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    good = {c: "" for c in HEAD.split(",")}
    good.update({"Order Date": d(-6), "Order Via": "Call", "Type of Submission": "Challan", "DC/Inv No": f"AN-{u}-1"})
    bad = {**good, "DC/Inv No": f"AN-{u}-2", "Order Via": "Nope"}
    res = client.post("/admin/bulk/orders/confirm", json={"rows": [good, bad]}, headers=admin).json()
    assert res["created"] == 0 and len(res["errors"]) == 1
    assert client.get(f"/admin/orders/search?q=AN-{u}", headers=admin).json()["total"] == 0


# ---------------------------------------------------------------- members
def test_bulk_members_validate_import_login_and_undo(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    head = "Display Name,Username,Role,Temporary Password"
    raw = (f"{head}\nAnil Kumar,anil.{u},godown,\nSeema Rao,seema.{u},shop,Temp-Pass-9999\n"
           f"Bad Admin,badadmin.{u},admin,\nBad Name,BAD NAME,shop,\nShort Pw,shortpw.{u},shop,abc\n"
           f"Weird Role,weird.{u},wizard,\nEXAMPLE Ravi,example.ravi,godown,\n").encode()
    rep = upload(client, admin, "users", raw)
    assert rep["total"] == 6 and rep["valid"] == 2
    errs = " | ".join(e["error"] for e in rep["errors"])
    for needle in ("cannot be bulk-created", "Username must be", "at least 10", "wizard"):
        assert needle in errs

    res = confirm(client, admin, "users", rep["rows"], "team.csv")
    creds = {c["username"]: c for c in res["credentials"]}
    assert creds[f"seema.{u}"]["temporary_password"] == "Temp-Pass-9999"
    generated = creds[f"anil.{u}"]["temporary_password"]
    assert len(generated) >= 10
    login_ok = client.post("/auth/login", json={"username": f"anil.{u}", "password": generated})
    assert login_ok.status_code == 200 and login_ok.json()["must_change_password"] is True
    assert client.post("/admin/bulk/users/confirm", json={"rows": rep["rows"]}, headers=admin).json()["created"] == 0

    # after someone really uses their account, the import cannot be undone any more
    tok = login_ok.json()["access_token"]
    ch = client.post("/auth/change-password", headers={"Authorization": f"Bearer {tok}"},
                     json={"current_password": generated, "new_password": "A-Brand-New-Pass-1"})
    assert ch.status_code == 200, ch.text
    blocked = client.post(f"/admin/bulk/batches/{res['batch_id']}/undo", headers=admin)
    assert blocked.status_code == 409


def test_bulk_members_undo_when_nobody_signed_in(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    rep = upload(client, admin, "users", f"Display Name,Username,Role\nUn Do,undo.{u},receiving\n".encode())
    res = confirm(client, admin, "users", rep["rows"])
    assert client.post(f"/admin/bulk/batches/{res['batch_id']}/undo", headers=admin).json()["removed"] == 1
    users = client.get("/admin/users", headers=admin).json()
    assert f"undo.{u}" not in [x["username"] for x in users]


# ---------------------------------------------------------------- filters
def _import(client, admin, lines):
    rep = upload(client, admin, "orders", order_csv(*lines))
    assert not rep["errors"], rep["errors"]
    return confirm(client, admin, "orders", rep["rows"])["batch_id"]


def test_filters_sorting_paging_and_stage_logic(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    batch = _import(client, admin, [
        f"{d(-30)},Call,Challan,F-{u}-1,Rohini,urgent,Ravi,Sonu,Delivered,Amit,{d(-30)} 14:30,50,{d(-29)},Paid,1500",
        f"{d(-25)},Walk-in,Invoice,F-{u}-2,Pitampura,,,,,,,,,,",
        f"{d(-20)},Online,Challan,F-{u}-3,Rohini,,,,Delivered,Amit,{d(-20)} 16:00,,,Pending,300",
        f"{d(-15)},Call,Challan,F-{u}-4,Dwarka,,,,Cancelled,,,,,,",
    ])
    base = f"/admin/orders/search?batch_id={batch}"

    def dcs(qs, **kw):
        r = client.get(base + qs, headers=admin).json()
        return [x["dc_inv_no"].split("-")[-1] for x in r["rows"]], r["total"]

    assert dcs("&sort=order_date&dir=asc")[0] == ["1", "2", "3", "4"]
    assert dcs("")[0] == ["4", "3", "2", "1"]                                   # newest first by default
    assert dcs("&stage=Closed")[0] == ["1"]
    assert dcs("&stage=Awaiting godown&stage=Cancelled&sort=sl_no&dir=asc")[0] == ["2", "4"]
    assert dcs("&stage=Awaiting receiving")[0] == ["3"]
    assert dcs("&channel=Call&sort=sl_no&dir=asc")[0] == ["1", "4"]
    assert dcs("&q=urgent")[0] == ["1"] and dcs("&q=rohini&sort=sl_no&dir=asc")[0] == ["1", "3"]
    lo, hi = (TODAY - timedelta(days=26)).isoformat(), (TODAY - timedelta(days=19)).isoformat()
    assert dcs(f"&date_from={lo}&date_to={hi}&sort=sl_no&dir=asc")[0] == ["2", "3"]
    assert dcs("&min_amount=1000")[0] == ["1"] and dcs("&max_amount=500")[0] == ["3"]
    assert dcs("&payment_status=__none__&sort=sl_no&dir=asc")[0] == ["2", "4"]  # "no payment status yet"
    assert dcs("&payment_status=__none__&payment_status=Pending&sort=sl_no&dir=asc")[0] == ["2", "3", "4"]
    assert dcs("&cancelled=exclude&sort=sl_no&dir=asc")[0] == ["1", "2", "3"]
    assert dcs("&cancelled=only")[0] == ["4"]
    assert dcs("&delivered_by=Amit&sort=sl_no&dir=asc")[0] == ["1", "3"]
    page = client.get(base + "&sort=sl_no&dir=asc&limit=2&offset=2", headers=admin).json()
    assert page["total"] == 4 and [x["dc_inv_no"][-1] for x in page["rows"]] == ["3", "4"]
    assert client.get(base + "&date_from=not-a-date", headers=admin).status_code == 422
    # an unknown sort column is ignored, not executed
    assert client.get(base + "&sort=DROP TABLE&dir=sideways", headers=admin).status_code == 200

    # archived orders are hidden unless asked for
    with db_conn.cursor() as cur:
        cur.execute("UPDATE fact_orders SET archived_at = now() WHERE dc_inv_no = %s", (f"F-{u}-1",))
    db_conn.commit()
    assert dcs("&sort=sl_no&dir=asc")[1] == 3
    with_arch = client.get(base + "&include_archived=1", headers=admin).json()
    assert with_arch["total"] == 4 and any(x["archived"] for x in with_arch["rows"])


def test_filter_options_and_sql_injection_is_inert(client, db_conn, seed):
    admin = login(client, db_conn)
    opts = client.get("/admin/orders/filter-options", headers=admin).json()
    assert "Call" in opts["channel"] and "Amit" in opts["delivered_by"] and "Closed" in opts["stage"]
    evil = client.get("/admin/orders/search", params={"q": "'; DROP TABLE fact_orders; --", "channel": "x' OR '1'='1"},
                      headers=admin)
    assert evil.status_code == 200 and evil.json()["total"] == 0
    assert client.get("/admin/orders/search?limit=5", headers=admin).status_code == 200


def test_saved_filters(client, db_conn, seed):
    a1, a2 = login(client, db_conn), login(client, db_conn)
    name = "Open " + uuid.uuid4().hex[:4]
    r = client.post("/admin/saved-filters", headers=a1,
                    json={"name": name, "shared": False, "definition": {"stage": ["Awaiting godown"], "junk": 1}})
    assert r.status_code == 201 and r.json()["definition"] == {"stage": ["Awaiting godown"]}
    assert name in [s["name"] for s in client.get("/admin/saved-filters", headers=a1).json()]
    assert name not in [s["name"] for s in client.get("/admin/saved-filters", headers=a2).json()]
    client.post("/admin/saved-filters", headers=a1, json={"name": name, "shared": True, "definition": {"q": "x"}})
    seen = [s for s in client.get("/admin/saved-filters", headers=a2).json() if s["name"] == name]
    assert len(seen) == 1 and seen[0]["mine"] is False and seen[0]["definition"] == {"q": "x"}
    assert client.delete(f"/admin/saved-filters/{seen[0]['filter_key']}", headers=a2).status_code == 404
    assert client.delete(f"/admin/saved-filters/{seen[0]['filter_key']}", headers=a1).status_code == 200
    assert client.post("/admin/saved-filters", headers=a1,
                       json={"name": "empty", "definition": {"junk": 1}}).status_code == 422


def test_filtered_excel_export(client, db_conn, seed):
    admin = login(client, db_conn)
    u = uuid.uuid4().hex[:5]
    batch = _import(client, admin, [f"{d(-5)},Call,Challan,X-{u}-1,=1+1,,,,,,,,,,",
                                    f"{d(-5)},Online,Invoice,X-{u}-2,,,,,,,,,,,"])
    r = client.get(f"/admin/orders/search.xlsx?batch_id={batch}&channel=Call", headers=admin)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb["Orders"].iter_rows(values_only=True))
    assert len(rows) == 2 and rows[1][1] == f"X-{u}-1"
    assert "channel" in [c[0] for c in wb["Filter"].iter_rows(values_only=True)]
    formula_cell = [c for c in rows[1] if c == "=1+1"]
    assert formula_cell, "an address starting with = must stay text"
