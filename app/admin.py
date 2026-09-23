"""Admin-only routes: members, dashboard numbers, Excel export, app links, audit log.
Also the /links route every logged-in user calls to see which apps they may open."""
import io
import os
import re
import secrets
from datetime import date, datetime, timedelta
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from psycopg2 import errors as pgerr
from psycopg2.extras import Json
from pydantic import BaseModel, ConfigDict, Field, field_validator

from . import auth, config, dashboard, db, roles

ADMIN = auth.require_roles("admin")
ANY_LOGGED_IN = auth.require_roles(*roles.ALL_ROLES)
router = APIRouter(prefix="/admin", tags=["admin"])
public_router = APIRouter(tags=["links"])

ASSIGNABLE_ROLES = sorted(roles.ALL_ROLES - {"legacy"})
USERNAME_RE = re.compile(r"^[a-z0-9][a-z0-9_.-]{2,39}$")
MAX_EXPORT_DAYS = 400


def audit(cur, admin, action, target=None, details=None):
    cur.execute("INSERT INTO admin_audit_log (admin_user_key, admin_username, action, target, details) "
                "VALUES (%s, %s, %s, %s, %s)",
                (admin["user_key"], admin["username"], action, target, Json(details) if details else None))


# ------------------------------------------------------------------ members
class NewUser(BaseModel):
    model_config = ConfigDict(extra="forbid")
    username: str
    display_name: str = Field(min_length=1, max_length=150)
    role: str
    password: str = Field(min_length=10, max_length=200)

    @field_validator("username")
    @classmethod
    def _u(cls, v):
        v = v.strip().lower()
        if not USERNAME_RE.match(v):
            raise ValueError("username: 3-40 chars, lowercase letters, digits, . _ -")
        return v

    @field_validator("display_name")
    @classmethod
    def _d(cls, v):
        v = " ".join(v.split())
        if not v:
            raise ValueError("display_name must not be blank")
        return v


class EditUser(BaseModel):
    model_config = ConfigDict(extra="forbid")
    display_name: Optional[str] = Field(default=None, min_length=1, max_length=150)
    role: Optional[str] = None


class PasswordIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    password: str = Field(min_length=10, max_length=200)


def _check_role(role):
    if role not in ASSIGNABLE_ROLES:
        raise HTTPException(422, f"Unknown role. Choose one of: {', '.join(ASSIGNABLE_ROLES)}")


@router.get("/roles")
def list_roles(admin=Depends(ADMIN)):
    return {"roles": ASSIGNABLE_ROLES, "order_roles": sorted(roles.ORDER_ROLES)}


@router.get("/users")
def list_users(admin=Depends(ADMIN)):
    with db.cursor() as cur:
        cur.execute("""
            SELECT u.user_key, u.username, u.display_name, u.role,
                   (u.password_hash = %s) AS disabled, u.must_change_password,
                   (SELECT count(*) FROM fact_orders o WHERE o.created_by_user_key = u.user_key) AS orders_created
            FROM dim_user u ORDER BY (u.password_hash = %s), u.role, u.username""",
                    (auth.DISABLED_HASH, auth.DISABLED_HASH))
        return cur.fetchall()


@router.post("/users", status_code=201)
def create_user(body: NewUser, admin=Depends(ADMIN)):
    _check_role(body.role)
    try:
        with db.cursor() as cur:
            cur.execute("INSERT INTO dim_user (username, password_hash, role, display_name, must_change_password) "
                        "VALUES (%s, %s, %s, %s, true) RETURNING user_key",
                        (body.username, auth.hash_password(body.password), body.role, body.display_name))
            key = cur.fetchone()["user_key"]
            audit(cur, admin, "user.create", body.username, {"role": body.role, "display_name": body.display_name})
    except pgerr.UniqueViolation:
        raise HTTPException(409, "That username already exists")
    return {"user_key": key}


def _get_user(cur, user_key):
    cur.execute("SELECT user_key, username, role, display_name FROM dim_user WHERE user_key = %s FOR UPDATE", (user_key,))
    row = cur.fetchone()
    if not row:
        raise HTTPException(404, "User not found")
    return row


@router.patch("/users/{user_key}")
def edit_user(user_key: int, body: EditUser, admin=Depends(ADMIN)):
    changes = body.model_dump(exclude_unset=True)
    if not changes:
        raise HTTPException(422, "Nothing to change")
    if "role" in changes:
        _check_role(changes["role"])
        if user_key == admin["user_key"] and changes["role"] != admin["role"]:
            raise HTTPException(400, "You cannot change your own role")
    if "display_name" in changes:
        changes["display_name"] = " ".join(changes["display_name"].split())
    with db.cursor() as cur:
        target = _get_user(cur, user_key)
        sets = ", ".join(f"{k} = %s" for k in changes)
        cur.execute(f"UPDATE dim_user SET {sets} WHERE user_key = %s", list(changes.values()) + [user_key])
        audit(cur, admin, "user.edit", target["username"],
              {"before": {k: target[k] for k in changes}, "after": changes})
    return {"ok": True}


@router.post("/users/{user_key}/reset-password")
def reset_password(user_key: int, body: PasswordIn, admin=Depends(ADMIN)):
    with db.cursor() as cur:
        target = _get_user(cur, user_key)
        cur.execute("UPDATE dim_user SET password_hash = %s, must_change_password = true WHERE user_key = %s",
                    (auth.hash_password(body.password), user_key))
        audit(cur, admin, "user.reset_password", target["username"])
    return {"ok": True}


@router.post("/users/{user_key}/disable")
def disable_user(user_key: int, admin=Depends(ADMIN)):
    if user_key == admin["user_key"]:
        raise HTTPException(400, "You cannot disable your own account")
    with db.cursor() as cur:
        target = _get_user(cur, user_key)
        cur.execute("UPDATE dim_user SET password_hash = %s WHERE user_key = %s", (auth.DISABLED_HASH, user_key))
        audit(cur, admin, "user.disable", target["username"])
    return {"ok": True}


# ------------------------------------------------------------------ dashboard
def _range(date_from, date_to):
    with db.cursor() as cur:
        today = dashboard.today_ist(cur)
    date_to = date_to or today
    date_from = date_from or (date_to - timedelta(days=29))
    if date_from > date_to:
        raise HTTPException(422, "date_from must not be after date_to")
    if (date_to - date_from).days > MAX_EXPORT_DAYS:
        raise HTTPException(422, f"Choose a range of at most {MAX_EXPORT_DAYS} days")
    return date_from, date_to


@router.get("/dashboard")
def get_dashboard(date_from: Optional[date] = None, date_to: Optional[date] = None, admin=Depends(ADMIN)):
    date_from, date_to = _range(date_from, date_to)
    return dashboard.build(date_from, date_to)


# Backs the dashboard's "click a bar to see the orders" drill-down. date_from/date_to
# scope it to the selected period (or omit both for the pipeline chart, which is
# always "right now" regardless of period); stage and the hour range narrow further.
@router.get("/orders")
def admin_list_orders(
    date_from: Optional[date] = None, date_to: Optional[date] = None,
    stage: Optional[str] = Query(default=None, description="one stage, or several comma-separated"),
    min_hours: Optional[float] = None, max_hours: Optional[float] = None,
    limit: int = Query(default=200, ge=1, le=500), admin=Depends(ADMIN),
):
    where, params = ["1=1"], []
    if date_from:
        where.append("order_date >= %s"); params.append(date_from)
    if date_to:
        where.append("order_date <= %s"); params.append(date_to)
    if stage:
        stages = [s.strip() for s in stage.split(",") if s.strip()]
        where.append("stage = ANY(%s)"); params.append(stages)
    if min_hours is not None:
        where.append("hours_to_deliver >= %s"); params.append(min_hours)
    if max_hours is not None:
        where.append("hours_to_deliver < %s"); params.append(max_hours)
    with db.cursor() as cur:
        cur.execute(f"""
            SELECT sl_no, dc_inv_no, order_date, stage, delivery_status, channel, submission_type,
                   shipping_location, amount_received, cartage, hours_to_deliver, created_by
            FROM v_orders WHERE {' AND '.join(where)}
            ORDER BY order_date DESC, sl_no DESC LIMIT %s""", params + [limit])
        return cur.fetchall()


# ------------------------------------------------------------------ Excel export
EXPORT_COLUMNS = [  # (header, view column, width)
    ("Sl No", "sl_no", 8), ("DC / Inv No", "dc_inv_no", 14), ("Order date", "order_date", 12),
    ("Order via", "channel", 12), ("Submission type", "submission_type", 24), ("Stage", "stage", 18),
    ("Delivery status", "delivery_status", 18), ("Ready by", "ready_by", 16),
    ("Colour making by", "colour_making_by", 16), ("Delivered by", "delivered_by", 22),
    ("Delivered-by phone", "delivered_by_phone", 15), ("Payment status", "payment_status", 20),
    ("Amount received (Rs)", "amount_received", 14), ("Cartage (Rs)", "cartage", 12),
    ("Hours to deliver", "hours_to_deliver", 12), ("Shipping location", "shipping_location", 30),
    ("Remarks", "detailed_remarks", 34), ("Logged at (IST)", "timestamp_created", 18),
    ("Material delivered (IST)", "material_delivery_datetime", 18), ("Date received", "date_of_receiving", 12),
    ("Created by", "created_by", 14), ("Last updated by", "last_updated_by", 14),
    ("Last updated (IST)", "last_updated_at", 18), ("Cancelled", "is_cancelled", 10),
]
IST_COLS = {"timestamp_created", "material_delivery_datetime", "last_updated_at"}


@router.get("/export.xlsx")
def export_xlsx(date_from: date, date_to: date, admin=Depends(ADMIN)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from zoneinfo import ZoneInfo

    date_from, date_to = _range(date_from, date_to)
    ist = ZoneInfo(dashboard.IST)
    with db.cursor() as cur:
        cur.execute(f"SELECT {', '.join(c for _, c, _ in EXPORT_COLUMNS)} FROM v_orders "
                    "WHERE order_date BETWEEN %s AND %s ORDER BY order_date, sl_no", (date_from, date_to))
        rows = cur.fetchall()
        audit(cur, admin, "export", f"{date_from} to {date_to}", {"rows": len(rows)})
    summary = dashboard.headline(date_from, date_to)

    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append([h for h, _, _ in EXPORT_COLUMNS])
    head_fill = PatternFill("solid", fgColor="10263D")
    for i, (h, _, w) in enumerate(EXPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in rows:
        line = []
        for _, col, _ in EXPORT_COLUMNS:
            v = r[col]
            if col in IST_COLS and isinstance(v, datetime):
                v = v.astimezone(ist).replace(tzinfo=None)
            elif col == "hours_to_deliver" and v is not None:
                v = round(float(v), 1)
            elif col in ("amount_received", "cartage") and v is not None:
                v = float(v)
            elif col == "is_cancelled":
                v = "Yes" if v else "No"
            line.append(v)
        ws.append(line)
        for c in ws[ws.max_row]:  # text that starts with "=" must stay text, never a formula
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
    fmt = {"order_date": "dd-mmm-yyyy", "date_of_receiving": "dd-mmm-yyyy",
           "timestamp_created": "dd-mmm-yy hh:mm", "material_delivery_datetime": "dd-mmm-yy hh:mm",
           "last_updated_at": "dd-mmm-yy hh:mm", "amount_received": "#,##0.00", "cartage": "#,##0.00",
           "hours_to_deliver": "0.0"}
    for i, (_, col, _) in enumerate(EXPORT_COLUMNS, 1):
        if col in fmt:
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2, max_row=ws.max_row):
                for c in cell:
                    c.number_format = fmt[col]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Summary")
    lines = [("Period", f"{date_from:%d-%b-%Y} to {date_to:%d-%b-%Y}"), ("Orders", summary["orders"]),
             ("Cancelled", summary["cancelled"]), ("Delivered", summary["delivered"]),
             ("Fully closed", summary["closed"]), ("Amount received (Rs)", float(summary["amount_received"])),
             ("Cartage (Rs)", float(summary["cartage"])),
             ("Avg hours to deliver", round(summary["avg_hours"], 1) if summary["avg_hours"] is not None else None),
             ("Median hours to deliver", round(summary["median_hours"], 1) if summary["median_hours"] is not None else None),
             ("Delivered within 24 h", round(summary["within_24h"], 3) if summary["within_24h"] is not None else None),
             ("Generated (IST)", datetime.now(ist).strftime("%d-%b-%Y %H:%M"))]
    for a, b in lines:
        s.append([a, b])
    s.column_dimensions["A"].width, s.column_dimensions["B"].width = 26, 28
    for c in s["A"]:
        c.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"vardhman_orders_{date_from}_to_{date_to}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ------------------------------------------------------------------ audit log
@router.get("/audit")
def get_audit(limit: int = Query(default=100, ge=1, le=500), admin=Depends(ADMIN)):
    with db.cursor() as cur:
        cur.execute("SELECT audit_key, at, admin_username, action, target, details FROM admin_audit_log "
                    "ORDER BY at DESC, audit_key DESC LIMIT %s", (limit,))
        return cur.fetchall()


# ------------------------------------------------------------------ security (client key)
@router.get("/security")
def get_security(admin=Depends(ADMIN)):
    has_db_value = bool(config.get(config.CLIENT_KEY_CONFIG_KEY))
    has_env_value = bool(os.environ.get("APPS_SCRIPT_CLIENT_KEY"))
    return {
        "client_key_configured": has_db_value or has_env_value,
        "client_key_source": "database" if has_db_value else ("environment" if has_env_value else "none"),
    }


@router.post("/security/client-key/rotate")
def rotate_client_key(admin=Depends(ADMIN)):
    """Generates a new secret and stores it in the database, replacing any env-var
    value. Returned once; admin must copy it into the one authorized Apps Script
    project's CLIENT_KEY Script Property, or every order-data call starts failing."""
    with db.cursor() as cur:
        new_key = secrets.token_urlsafe(48)
        config.set(config.CLIENT_KEY_CONFIG_KEY, new_key, updated_by=admin["username"])
        audit(cur, admin, "security.rotate_client_key")
    return {"client_key": new_key}


# ------------------------------------------------------------------ permissions (field access per role)
FIELD_NAMES = sorted({
    "order_received_date", "order_via_key", "submission_type_key", "dc_inv_no",
    "shipping_location", "detailed_remarks", "ready_by_person_key", "colour_making_person_key",
    "delivery_status_key", "material_delivery_datetime", "delivered_by_person_key", "cartage",
    "date_of_receiving", "payment_status_key", "amount_received",
})


@router.get("/permissions")
def get_permissions(admin=Depends(ADMIN)):
    rows = config.all_field_permissions()
    have = {(r["role"], r["field_name"]) for r in rows if r["editable"]}
    return {
        "roles": sorted(roles.ORDER_ROLES),
        "fields": FIELD_NAMES,
        "editable": [{"role": role, "field_name": f, "editable": (role, f) in have}
                     for role in sorted(roles.ORDER_ROLES) for f in FIELD_NAMES],
    }


class PermissionIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    role: str
    field_name: str
    editable: bool

    @field_validator("role")
    @classmethod
    def _role(cls, v):
        if v not in roles.ORDER_ROLES:
            raise ValueError(f"role must be one of: {', '.join(sorted(roles.ORDER_ROLES))}")
        return v

    @field_validator("field_name")
    @classmethod
    def _field(cls, v):
        if v not in FIELD_NAMES:
            raise ValueError(f"unknown field: {v}")
        return v


@router.put("/permissions")
def set_permission(body: PermissionIn, admin=Depends(ADMIN)):
    with db.cursor() as cur:
        config.set_field_permission(body.role, body.field_name, body.editable)
        audit(cur, admin, "permission.set", f"{body.role}.{body.field_name}", {"editable": body.editable})
    return {"ok": True}


# ------------------------------------------------------------------ app links
class LinkIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = Field(min_length=1, max_length=100)
    url: str = Field(max_length=2000)
    roles: List[str] = []
    sort_order: int = 0
    active: bool = True

    @field_validator("url")
    @classmethod
    def _https(cls, v):
        v = v.strip()
        if not v.startswith("https://"):
            raise ValueError("url must start with https://")
        return v

    @field_validator("roles")
    @classmethod
    def _roles(cls, v):
        bad = [r for r in v if r not in ASSIGNABLE_ROLES]
        if bad:
            raise ValueError(f"unknown role(s): {', '.join(bad)}")
        return sorted(set(v))


@router.get("/links")
def admin_links(admin=Depends(ADMIN)):
    with db.cursor() as cur:
        cur.execute("SELECT link_key, name, url, roles, sort_order, active FROM app_links ORDER BY sort_order, name")
        return cur.fetchall()


@router.post("/links", status_code=201)
def add_link(body: LinkIn, admin=Depends(ADMIN)):
    with db.cursor() as cur:
        cur.execute("INSERT INTO app_links (name, url, roles, sort_order, active) VALUES (%s, %s, %s, %s, %s) "
                    "RETURNING link_key", (body.name, body.url, body.roles, body.sort_order, body.active))
        key = cur.fetchone()["link_key"]
        audit(cur, admin, "link.create", body.name, body.model_dump())
    return {"link_key": key}


@router.put("/links/{link_key}")
def edit_link(link_key: int, body: LinkIn, admin=Depends(ADMIN)):
    with db.cursor() as cur:
        cur.execute("UPDATE app_links SET name=%s, url=%s, roles=%s, sort_order=%s, active=%s WHERE link_key=%s",
                    (body.name, body.url, body.roles, body.sort_order, body.active, link_key))
        if cur.rowcount == 0:
            raise HTTPException(404, "Link not found")
        audit(cur, admin, "link.edit", body.name, body.model_dump())
    return {"ok": True}


@router.delete("/links/{link_key}")
def delete_link(link_key: int, admin=Depends(ADMIN)):
    with db.cursor() as cur:
        cur.execute("DELETE FROM app_links WHERE link_key = %s RETURNING name", (link_key,))
        row = cur.fetchone()
        if not row:
            raise HTTPException(404, "Link not found")
        audit(cur, admin, "link.delete", row["name"])
    return {"ok": True}


# ------------------------------------------------------------------ what each user may open
@public_router.get("/links")
def my_links(user=Depends(ANY_LOGGED_IN)):
    with db.cursor() as cur:
        if user["role"] == "admin":
            cur.execute("SELECT link_key, name, url FROM app_links WHERE active ORDER BY sort_order, name")
        else:
            cur.execute("SELECT link_key, name, url FROM app_links WHERE active AND %s = ANY(roles) "
                        "ORDER BY sort_order, name", (user["role"],))
        return cur.fetchall()
