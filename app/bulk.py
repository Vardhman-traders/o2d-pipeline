"""Generic bulk-upload plumbing: read a CSV/Excel file, check it looks like the template, hand rows to an entity's
validator, and build the downloadable template. Same flow as OmniFlow's bulk upload:

  download template (with EXAMPLE rows) -> upload -> validate -> fix errored rows inline -> import -> undo

Nothing here touches the database; the per-entity rules live in bulk_entities.py.
"""
import csv
import io
import zipfile
from datetime import date, datetime

from fastapi import HTTPException

MAX_ROWS = 2000
MAX_BYTES = 5 * 1024 * 1024
EXAMPLE_MARK = "EXAMPLE"


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_xlsx(raw: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb["Template"] if "Template" in wb.sheetnames else wb.worksheets[0]   # our template has other sheets too
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if not header:
        return [], []
    names = [_cell(h) for h in header]
    rows = [{n: _cell(v) for n, v in zip(names, r, strict=False) if n} for r in it]
    return rows, [n for n in names if n]


def _read_csv(raw: bytes):
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252")  # Excel's "CSV" on Windows
    sample = text[:4096]
    delim = max((",", ";", "\t"), key=sample.count)
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    names = [(n or "").strip() for n in (reader.fieldnames or []) if n]
    rows = [{(k or "").strip(): _cell(v) for k, v in r.items() if k} for r in reader]
    return rows, names


def read_table(raw: bytes) -> tuple[list[dict], list[str]]:
    """(rows, header names). Rows are dicts of stripped strings; each gets `_row` = its line number in the file."""
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "File is too large (limit 5 MB).")
    if not raw.strip():
        return [], []
    try:
        rows, names = _read_xlsx(raw) if zipfile.is_zipfile(io.BytesIO(raw)) else _read_csv(raw)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Could not read this file. Please upload the CSV template (or an .xlsx made from it).")
    kept = []
    for i, r in enumerate(rows, start=2):
        if not any(v for v in r.values()):
            continue                                                   # blank line
        if any(str(v).upper().startswith(EXAMPLE_MARK) for v in r.values()):
            continue                                                   # the template's example rows
        r["_row"] = i
        kept.append(r)
    return kept, names


def header_problem(names: list[str], required: list[str], expected: list[str]) -> str | None:
    """A message when the file clearly is not the template (none of the required columns present)."""
    have = {n.lower() for n in names}
    missing = [c for c in required if c.lower() not in have]
    if required and len(missing) == len(required):
        return ("This file doesn't match the template. Expected columns such as: " + ", ".join(expected) +
                ". Found: " + (", ".join(names) if names else "(nothing - the file may be empty)") + ".")
    if missing:
        return "Missing required column(s): " + ", ".join(missing) + ". Please use the template."
    return None


def canonical_rows(rows: list[dict], expected: list[str]) -> list[dict]:
    """Re-key each row by the template's exact header spelling (matching case-insensitively)."""
    lookup = {c.lower(): c for c in expected}
    out = []
    for r in rows:
        row = {"_row": r.get("_row")}
        for k, v in r.items():
            if k == "_row":
                continue
            row[lookup.get(k.lower(), k)] = v
        for c in expected:
            row.setdefault(c, "")
        out.append(row)
    return out


def template_csv(columns: list[str], examples: list[list[str]]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    for ex in examples:
        w.writerow(ex)
    return "﻿" + buf.getvalue()  # BOM so Excel opens it as UTF-8


def check_row_cap(n: int) -> None:
    if n > MAX_ROWS:
        raise HTTPException(400, f"Too many rows - the maximum per upload is {MAX_ROWS}.")
